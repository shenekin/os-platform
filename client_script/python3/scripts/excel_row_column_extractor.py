#!/usr/bin/env python3
"""Extract configured rows and columns from an Excel workbook."""

from __future__ import annotations

import argparse
from copy import copy
import json
import re
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError as error:
    if error.name != "openpyxl":
        raise
    print(
        "Missing dependency: openpyxl. Run "
        "'.venv/bin/python -m pip install -r requirements.txt' "
        "and execute this script with '.venv/bin/python'.",
        file=sys.stderr,
    )
    raise SystemExit(1)


def load_config(config_path: Path) -> dict[str, Any]:
    with config_path.open("r", encoding="utf-8") as config_file:
        config = json.load(config_file)
    if not isinstance(config, dict):
        raise ValueError("The configuration root must be a JSON object.")
    return config


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    cell_value = getattr(value, "value", value)
    return "" if cell_value is None else str(cell_value).strip()


def column_index(column: str | int, headers: list[str]) -> int:
    if isinstance(column, int):
        if column < 1:
            raise ValueError(f"Column indexes must start at 1: {column}")
        return column

    column_text = str(column).strip()
    if column_text.isdigit():
        return column_index(int(column_text), headers)

    if re.fullmatch(r"[A-Za-z]+", column_text):
        result = 0
        for character in column_text.upper():
            result = result * 26 + ord(character) - ord("A") + 1
        return result

    matching_indexes = [
        index for index, header in enumerate(headers, start=1) if header == column_text
    ]
    if not matching_indexes:
        raise ValueError(f"Column was not found by header name: {column_text}")
    return matching_indexes[0]


def compare(value: str, operator: str, expected: str) -> bool:
    if operator == "not_empty":
        return bool(value)
    if operator == "empty":
        return not value
    if operator == "equals":
        return value == expected
    if operator == "not_equals":
        return value != expected
    if operator == "contains":
        return expected in value
    if operator == "not_contains":
        return expected not in value
    if operator == "starts_with":
        return value.startswith(expected)
    if operator == "ends_with":
        return value.endswith(expected)
    if operator == "regex":
        return re.search(expected, value) is not None
    raise ValueError(f"Unsupported condition operator: {operator}")


def row_matches(
    row: tuple[Any, ...],
    headers: list[str],
    conditions: list[dict[str, Any]],
    mode: str,
) -> bool:
    results = []
    for condition in conditions:
        if "column" not in condition or "operator" not in condition:
            raise ValueError("Each row condition needs a column and an operator.")
        index = column_index(condition["column"], headers)
        value = cell_text(row[index - 1] if index <= len(row) else None)
        expected = cell_text(condition.get("value"))
        results.append(compare(value, condition["operator"], expected))

    if not results:
        return True
    if mode == "all":
        return all(results)
    if mode == "any":
        return any(results)
    raise ValueError("condition_mode must be either 'all' or 'any'.")


def select_columns(columns: list[Any], headers: list[str]) -> list[int]:
    if not columns:
        return list(range(1, len(headers) + 1))
    indexes = [column_index(column, headers) for column in columns]
    if len(set(indexes)) != len(indexes):
        raise ValueError("output_columns contains duplicate columns.")
    return indexes


def preserve_format_indexes(
    format_values: list[dict[str, Any]], headers: list[str]
) -> dict[int, set[str]]:
    result: dict[int, set[str]] = {}
    for item in format_values:
        if "column" not in item or "values" not in item:
            raise ValueError("Each preserve_format_values item needs column and values.")
        values = item["values"]
        if not isinstance(values, list):
            raise ValueError("preserve_format_values values must be a list.")
        index = column_index(item["column"], headers)
        result[index] = {cell_text(value) for value in values}
    return result


def format_column_indexes(columns: list[Any], headers: list[str]) -> set[int]:
    return {column_index(column, headers) for column in columns}


def resolve_config_path(path_value: str, config_directory: Path) -> Path:
    path = Path(path_value).expanduser()
    if path.is_absolute():
        return path
    return (config_directory / path).resolve()


def extract(
    config: dict[str, Any], config_directory: Path | None = None
) -> tuple[int, int, Path]:
    base_directory = config_directory or Path.cwd()
    input_path = resolve_config_path(config["input_file"], base_directory)
    output_path = resolve_config_path(config["output_file"], base_directory)
    sheet_name = config.get("sheet_name")
    header_row = int(config.get("header_row", 1))
    data_start_row = int(config.get("data_start_row", header_row + 1))
    if header_row < 1:
        raise ValueError("header_row must be at least 1.")
    if data_start_row <= header_row:
        raise ValueError("data_start_row must be greater than header_row.")

    workbook = load_workbook(
        input_path,
        read_only=False,
        data_only=bool(config.get("data_only", False)),
    )
    try:
        headers: list[str] | None = None
        output_indexes: list[int] | None = None
        output_widths: list[float | None] | None = None
        selected_rows: list[tuple[list[Any], list[Any], float | None]] = []
        matched_rows = 0
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet was not found: {sheet_name}")
            worksheets = [workbook[sheet_name]]
        else:
            worksheets = workbook.worksheets

        for worksheet in worksheets:
            sheet_headers: list[str] | None = None
            sheet_output_indexes: list[int] | None = None
            sheet_preserve_format_indexes: dict[int, set[str]] = {}
            sheet_preserve_format_columns: set[int] = set()
            for row_number, row in enumerate(worksheet.iter_rows(values_only=False), start=1):
                if row_number < header_row:
                    continue
                if row_number == header_row:
                    sheet_headers = [cell_text(value) for value in row]
                    sheet_output_indexes = select_columns(
                        config.get("output_columns", []), sheet_headers
                    )
                    sheet_preserve_format_indexes = preserve_format_indexes(
                        config.get("preserve_format_values", []), sheet_headers
                    )
                    sheet_preserve_format_columns = format_column_indexes(
                        config.get("preserve_format_columns", []), sheet_headers
                    )
                    if headers is None:
                        headers = sheet_headers
                        output_indexes = sheet_output_indexes
                        output_widths = [
                            worksheet.column_dimensions[get_column_letter(index)].width
                            for index in sheet_output_indexes
                        ]
                    continue
                if row_number < data_start_row:
                    continue
                if sheet_headers is None or sheet_output_indexes is None:
                    raise ValueError(
                        f"Worksheet '{worksheet.title}' does not contain the configured header row."
                    )
                if row_matches(
                    row,
                    sheet_headers,
                    config.get("row_conditions", []),
                    config.get("condition_mode", "all"),
                ):
                    matched_rows += 1
                    values = [
                        row[index - 1].value if index <= len(row) else None
                        for index in sheet_output_indexes
                    ]
                    style_sources = []
                    for index in sheet_output_indexes:
                        source_cell = row[index - 1] if index <= len(row) else None
                        if (
                            source_cell is not None
                            and (
                                index in sheet_preserve_format_columns
                                or (
                                    index in sheet_preserve_format_indexes
                                    and cell_text(source_cell.value)
                                    in sheet_preserve_format_indexes[index]
                                )
                            )
                        ):
                            style_sources.append(source_cell)
                        else:
                            style_sources.append(None)
                    selected_rows.append(
                        (values, style_sources, worksheet.row_dimensions[row_number].height)
                    )

        if headers is None or output_indexes is None or output_widths is None:
            raise ValueError("The worksheet does not contain the configured header row.")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        result = Workbook()
        result_sheet = result.active
        result_sheet.title = config.get("output_sheet_name", "Extracted Data")
        for output_column, width in enumerate(output_widths, start=1):
            if width is not None:
                result_sheet.column_dimensions[get_column_letter(output_column)].width = width
        result_sheet.append([headers[index - 1] for index in output_indexes])
        for selected_values, style_sources, source_row_height in selected_rows:
            result_sheet.append(selected_values)
            output_row = result_sheet.max_row
            if source_row_height is not None:
                result_sheet.row_dimensions[output_row].height = source_row_height
            for output_column, source_cell in enumerate(style_sources, start=1):
                if source_cell is not None:
                    result_cell = result_sheet.cell(output_row, output_column)
                    result_cell.font = copy(source_cell.font)
                    result_cell.fill = copy(source_cell.fill)
                    result_cell.border = copy(source_cell.border)
                    result_cell.alignment = copy(source_cell.alignment)
                    result_cell.protection = copy(source_cell.protection)
                    result_cell.number_format = source_cell.number_format
        result.save(output_path)
        return matched_rows, len(output_indexes), output_path
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extract configured rows and columns from an Excel workbook."
    )
    parser.add_argument(
        "--config",
        default="config/config.json",
        help="Path to the JSON configuration file.",
    )
    args = parser.parse_args()
    try:
        config_path = Path(args.config).expanduser().resolve()
        matched_rows, selected_columns, output_path = extract(
            load_config(config_path), config_path.parent
        )
    except (KeyError, OSError, ValueError, json.JSONDecodeError) as error:
        print(f"Extraction failed: {error}", file=sys.stderr)
        return 1
    print(
        f"Extraction completed: {matched_rows} matching rows and "
        f"{selected_columns} columns written to {output_path}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())